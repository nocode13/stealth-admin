"""Прайс: код товара → строки со всех листов книги.

Листы устроены по-разному (в «сан» колонки «оптом, розница», в «Лист4» наоборот, в «Лист1»
розницы нет, зато есть строки-разделители «Склад: …»), поэтому колонки ищем по заголовку.
"""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook

NUMBER_RE = re.compile(r"-?\d+(?:[.,]\d+)?")
HEADER_SCAN_ROWS = 10


@dataclass(frozen=True)
class ExcelRow:
    sheet: str
    warehouse: str | None
    name: str
    stock: float | None
    wholesale_usd: float | None
    retail_usd: float | None

    def to_json(self) -> dict:
        data = asdict(self)
        return {
            "sheets": [data["sheet"]],
            "warehouse": data["warehouse"],
            "name": data["name"],
            "stock": data["stock"],
            "wholesaleUsd": data["wholesale_usd"],
            "retailUsd": data["retail_usd"],
        }


def to_number(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = NUMBER_RE.search(str(value).replace("\xa0", "").replace(" ", ""))
    return float(match.group(0).replace(",", ".")) if match else None


def to_price(value: object) -> float | None:
    """0 и пусто в прайсе значат одно — «цены нет»."""
    number = to_number(value)
    return number if number and number > 0 else None


def normalize_code(value: object) -> str | None:
    """Коды в прайсе — 8 цифр с ведущими нулями; если Excel сделал из них число, дополняем."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text.zfill(8) if text.isdigit() else None


def _column(headers: tuple, *patterns: str) -> int | None:
    for index, header in enumerate(headers):
        if header is None:
            continue
        text = str(header).strip().lower()
        if any(re.search(pattern, text) for pattern in patterns):
            return index
    return None


def _cell(row: tuple, index: int | None) -> object:
    return row[index] if index is not None and index < len(row) else None


def load_price(path: Path) -> dict[str, list[ExcelRow]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    index: dict[str, list[ExcelRow]] = {}

    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers: tuple | None = None
        for _ in range(HEADER_SCAN_ROWS):
            row = next(rows, None)
            if row is None:
                break
            if _column(row, r"^код товара") is not None:
                headers = row
                break
        if headers is None:
            continue

        code_col = _column(headers, r"^код товара")
        name_col = _column(headers, r"^наименование")
        stock_col = _column(headers, r"на складе")
        retail_col = _column(headers, r"розни")
        wholesale_col = _column(headers, r"опт")
        warehouse: str | None = None

        for row in rows:
            first = _cell(row, code_col)
            if isinstance(first, str) and first.strip().lower().startswith("склад"):
                name = _cell(row, name_col)
                warehouse = str(name).strip() if name else None
                continue
            code = normalize_code(first)
            name = _cell(row, name_col)
            if not code or not name:
                continue
            index.setdefault(code, []).append(
                ExcelRow(
                    sheet=sheet.title,
                    warehouse=warehouse,
                    name=" ".join(str(name).split()),
                    stock=to_number(_cell(row, stock_col)),
                    wholesale_usd=to_price(_cell(row, wholesale_col)),
                    retail_usd=to_price(_cell(row, retail_col)),
                )
            )

    workbook.close()
    return index


def collapse_rows(rows: list[ExcelRow]) -> tuple[list[dict], bool]:
    """Одинаковые строки с разных листов (копия «Лист2 (2)» = «сан») схлопываем в одну.

    Конфликт — когда по одному коду расходятся розница или название: это показываем на карточке.
    Разный остаток по складам конфликтом не считаем, это нормально.
    """
    merged: dict[tuple, dict] = {}
    for row in rows:
        key = (row.name, row.warehouse, row.stock, row.wholesale_usd, row.retail_usd)
        if key in merged:
            if row.sheet not in merged[key]["sheets"]:
                merged[key]["sheets"].append(row.sheet)
        else:
            merged[key] = row.to_json()

    result = list(merged.values())
    # Строки с розницей — первыми: именно их карточка берёт по умолчанию.
    result.sort(key=lambda r: r["retailUsd"] is None)
    retails = {r["retailUsd"] for r in result if r["retailUsd"] is not None}
    names = {r["name"].lower() for r in result}
    return result, len(retails) > 1 or len(names) > 1
